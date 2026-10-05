import sqlite3
import json
import datetime
import os
import sys
from openpyxl import Workbook, load_workbook
from config import BASE_DIR

if hasattr(sys.stdout, 'reconfigure'):
    try: sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception: pass
if hasattr(sys.stderr, 'reconfigure'):
    try: sys.stderr.reconfigure(encoding='utf-8', errors='replace')
    except Exception: pass

DB_NAME = BASE_DIR / 'orders.db'
EXCEL_NAME = BASE_DIR / 'orders_record.xlsx'

def get_db_connection():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()
    # 建立基礎表格
    cursor.execute('''
    CREATE TABLE IF NOT EXISTS orders (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        guild_id TEXT,
        order_num TEXT UNIQUE,
        customer_id TEXT,
        type TEXT,
        price REAL,
        players_needed INTEGER,
        remarks TEXT,
        status TEXT,
        admin_id TEXT,
        payment_method TEXT,
        player_ids TEXT DEFAULT '[]',
        ticket_msg_id TEXT,
        ticket_channel_id TEXT,
        dispatch_msg_id TEXT,
        dispatch_channel_id TEXT,
        review_submitted INTEGER DEFAULT 0,
        created_at TEXT,
        payment_time TEXT,
        completed_time TEXT
    )
    ''')
    
    cursor.execute('''
    CREATE TABLE IF NOT EXISTS spending_adjustments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        guild_id TEXT,
        customer_id TEXT,
        amount REAL,
        reason TEXT,
        admin_id TEXT,
        created_at TEXT
    )
    ''')

    cursor.execute('''
    CREATE TABLE IF NOT EXISTS stock_shipments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        guild_id TEXT,
        product_name TEXT NOT NULL,
        channel_id TEXT NOT NULL,
        message_id TEXT NOT NULL,
        line_number INTEGER NOT NULL,
        account TEXT NOT NULL,
        password TEXT NOT NULL,
        shipped_by TEXT NOT NULL,
        shipped_at TEXT NOT NULL,
        UNIQUE(channel_id, message_id, line_number)
    )
    ''')

    cursor.execute('''
    CREATE TABLE IF NOT EXISTS guild_settings (
        guild_id TEXT PRIMARY KEY,
        dispatch_channel_id TEXT,
        log_channel_id TEXT,
        review_channel_id TEXT,
        stock_100m_channel_id TEXT,
        stock_200m_channel_id TEXT,
        shipping_log_channel_id TEXT,
        order_disclaimer TEXT
    )
    ''')

    cursor.execute('''
    CREATE TABLE IF NOT EXISTS guild_stock_products (
        guild_id TEXT NOT NULL,
        product_name TEXT NOT NULL,
        channel_id TEXT NOT NULL,
        PRIMARY KEY (guild_id, product_name)
    )
    ''')
    cursor.execute('''
    CREATE TABLE IF NOT EXISTS guild_payment_accounts (
        guild_id TEXT NOT NULL,
        payment_method TEXT NOT NULL,
        account_info TEXT NOT NULL,
        PRIMARY KEY (guild_id, payment_method)
    )
    ''')
    # --- 自動修復：檢查並補上缺少的欄位 ---
    cursor.execute("PRAGMA table_info(orders)")
    columns = [column[1] for column in cursor.fetchall()]
    
    if 'payment_time' not in columns:
        cursor.execute('ALTER TABLE orders ADD COLUMN payment_time TEXT')
        print("✅ 已自動補上 payment_time 欄位")
        
    if 'completed_time' not in columns:
        cursor.execute('ALTER TABLE orders ADD COLUMN completed_time TEXT')
        print("✅ 已自動補上 completed_time 欄位")

    if 'player_ids' not in columns:
        cursor.execute("ALTER TABLE orders ADD COLUMN player_ids TEXT DEFAULT '[]'")
        print("✅ 已自動補上 player_ids 欄位")

    if 'dispatch_msg_id' not in columns:
        cursor.execute('ALTER TABLE orders ADD COLUMN dispatch_msg_id TEXT')
        print("✅ 已自動補上 dispatch_msg_id 欄位")

    if 'ticket_msg_id' not in columns:
        cursor.execute('ALTER TABLE orders ADD COLUMN ticket_msg_id TEXT')
        print("✅ 已自動補上 ticket_msg_id 欄位")

    if 'ticket_channel_id' not in columns:
        cursor.execute('ALTER TABLE orders ADD COLUMN ticket_channel_id TEXT')
        print("✅ 已自動補上 ticket_channel_id 欄位")

    if 'dispatch_channel_id' not in columns:
        cursor.execute('ALTER TABLE orders ADD COLUMN dispatch_channel_id TEXT')
        print("✅ 已自動補上 dispatch_channel_id 欄位")

    if 'review_submitted' not in columns:
        cursor.execute('ALTER TABLE orders ADD COLUMN review_submitted INTEGER DEFAULT 0')
        print("✅ 已自動補上 review_submitted 欄位")

    if 'guild_id' not in columns:
        cursor.execute('ALTER TABLE orders ADD COLUMN guild_id TEXT')
        print("✅ 已自動補上 orders.guild_id 欄位")

    for table in ('spending_adjustments', 'stock_shipments'):
        cursor.execute(f"PRAGMA table_info({table})")
        table_columns = [column[1] for column in cursor.fetchall()]
        if 'guild_id' not in table_columns:
            cursor.execute(f'ALTER TABLE {table} ADD COLUMN guild_id TEXT')
            print(f"✅ 已自動補上 {table}.guild_id 欄位")

    cursor.execute("PRAGMA table_info(guild_settings)")
    settings_columns = [column[1] for column in cursor.fetchall()]
    if 'review_channel_id' not in settings_columns:
        cursor.execute('ALTER TABLE guild_settings ADD COLUMN review_channel_id TEXT')
        print("✅ 已自動補上 guild_settings.review_channel_id 欄位")

    settings_columns = [
        column[1] for column in cursor.execute(
            "PRAGMA table_info(guild_settings)"
        ).fetchall()
    ]
    for column_name in (
        'dispatch_channel_id',
        'log_channel_id',
        'stock_100m_channel_id',
        'stock_200m_channel_id',
        'shipping_log_channel_id',
        'order_disclaimer',
    ):
        if column_name not in settings_columns:
            cursor.execute(
                f'ALTER TABLE guild_settings ADD COLUMN {column_name} TEXT'
            )
            print(f"✅ 已自動補上 guild_settings.{column_name} 欄位")

    # 舊版固定商品設定要在所有欄位 migration 完成後再回填。
    cursor.execute('''
        INSERT OR IGNORE INTO guild_stock_products (guild_id, product_name, channel_id)
        SELECT guild_id, '100m幣號', stock_100m_channel_id
        FROM guild_settings
        WHERE stock_100m_channel_id IS NOT NULL
    ''')
    cursor.execute('''
        INSERT OR IGNORE INTO guild_stock_products (guild_id, product_name, channel_id)
        SELECT guild_id, '200m幣號', stock_200m_channel_id
        FROM guild_settings
        WHERE stock_200m_channel_id IS NOT NULL
    ''')
        
    conn.commit()
    conn.close()

def get_guild_settings(guild_id):
    conn = get_db_connection()
    settings = conn.execute(
        'SELECT * FROM guild_settings WHERE guild_id = ?', (str(guild_id),)
    ).fetchone()
    conn.close()
    return settings

def set_guild_order_disclaimer(guild_id, disclaimer):
    conn = get_db_connection()
    conn.execute(
        '''INSERT INTO guild_settings (guild_id, order_disclaimer)
           VALUES (?, ?)
           ON CONFLICT(guild_id) DO UPDATE SET
               order_disclaimer = excluded.order_disclaimer''',
        (str(guild_id), disclaimer),
    )
    conn.commit()
    conn.close()

def set_guild_settings(guild_id, dispatch_channel_id=None, log_channel_id=None,
                        review_channel_id=None,
                        stock_100m_channel_id=None, stock_200m_channel_id=None,
                        shipping_log_channel_id=None):
    conn = get_db_connection()
    conn.execute('''
        INSERT INTO guild_settings (
            guild_id, dispatch_channel_id, log_channel_id, review_channel_id,
            stock_100m_channel_id, stock_200m_channel_id, shipping_log_channel_id
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(guild_id) DO UPDATE SET
            dispatch_channel_id = excluded.dispatch_channel_id,
            log_channel_id = excluded.log_channel_id,
            review_channel_id = excluded.review_channel_id,
            stock_100m_channel_id = excluded.stock_100m_channel_id,
            stock_200m_channel_id = excluded.stock_200m_channel_id,
            shipping_log_channel_id = excluded.shipping_log_channel_id
    ''', (
        str(guild_id),
        str(dispatch_channel_id) if dispatch_channel_id else None,
        str(log_channel_id) if log_channel_id else None,
        str(review_channel_id) if review_channel_id else None,
        str(stock_100m_channel_id) if stock_100m_channel_id else None,
        str(stock_200m_channel_id) if stock_200m_channel_id else None,
        str(shipping_log_channel_id) if shipping_log_channel_id else None,
    ))
    conn.commit()
    conn.close()

def set_guild_channel(guild_id, purpose, channel_id):
    allowed = {
        "dispatch": "dispatch_channel_id",
        "log": "log_channel_id",
        "100m": "stock_100m_channel_id",
        "200m": "stock_200m_channel_id",
        "shipping_log": "shipping_log_channel_id",
    }
    column = allowed.get(purpose)
    if column is None:
        raise ValueError(f"不支援的頻道用途：{purpose}")

    conn = get_db_connection()
    conn.execute(
        f'''INSERT INTO guild_settings (guild_id, {column})
            VALUES (?, ?)
            ON CONFLICT(guild_id) DO UPDATE SET {column} = excluded.{column}''',
        (str(guild_id), str(channel_id)),
    )
    conn.commit()
    conn.close()

def clear_guild_channel(guild_id, purpose):
    allowed = {
        "dispatch": "dispatch_channel_id",
        "log": "log_channel_id",
        "100m": "stock_100m_channel_id",
        "200m": "stock_200m_channel_id",
        "shipping_log": "shipping_log_channel_id",
    }
    column = allowed.get(purpose)
    if column is None:
        raise ValueError(f"不支援的頻道用途：{purpose}")

    conn = get_db_connection()
    conn.execute(
        f'UPDATE guild_settings SET {column} = NULL WHERE guild_id = ?',
        (str(guild_id),),
    )
    conn.commit()
    conn.close()

def get_guild_products(guild_id):
    conn = get_db_connection()
    rows = conn.execute(
        'SELECT product_name, channel_id FROM guild_stock_products '
        'WHERE guild_id = ? ORDER BY product_name',
        (str(guild_id),),
    ).fetchall()
    conn.close()
    return rows

def get_guild_product_channel(guild_id, product_name):
    conn = get_db_connection()
    row = conn.execute(
        'SELECT channel_id FROM guild_stock_products '
        'WHERE guild_id = ? AND product_name = ?',
        (str(guild_id), product_name),
    ).fetchone()
    conn.close()
    return row['channel_id'] if row else None

def set_guild_product(guild_id, product_name, channel_id):
    conn = get_db_connection()
    existing = conn.execute(
        '''SELECT product_name FROM guild_stock_products
           WHERE guild_id = ? AND channel_id = ? AND product_name != ?''',
        (str(guild_id), str(channel_id), product_name),
    ).fetchone()
    if existing:
        conn.close()
        raise ValueError(
            f"庫存頻道已被商品 `{existing['product_name']}` 使用，"
            "不同商品不可共用同一庫存頻道。"
        )
    conn.execute('''
        INSERT INTO guild_stock_products (guild_id, product_name, channel_id)
        VALUES (?, ?, ?)
        ON CONFLICT(guild_id, product_name) DO UPDATE SET
            channel_id = excluded.channel_id
    ''', (str(guild_id), product_name, str(channel_id)))
    conn.commit()
    conn.close()

def delete_guild_product(guild_id, product_name):
    conn = get_db_connection()
    conn.execute(
        'DELETE FROM guild_stock_products WHERE guild_id = ? AND product_name = ?',
        (str(guild_id), product_name),
    )
    deleted = conn.total_changes
    conn.commit()
    conn.close()
    return deleted > 0

def get_guild_payment_account(guild_id, payment_method):
    conn = get_db_connection()
    row = conn.execute(
        'SELECT account_info FROM guild_payment_accounts '
        'WHERE guild_id = ? AND payment_method = ?',
        (str(guild_id), payment_method),
    ).fetchone()
    conn.close()
    return row['account_info'] if row else None

def get_guild_payment_methods(guild_id):
    conn = get_db_connection()
    rows = conn.execute(
        'SELECT payment_method, account_info FROM guild_payment_accounts '
        'WHERE guild_id = ? ORDER BY payment_method',
        (str(guild_id),),
    ).fetchall()
    conn.close()
    return rows

def set_guild_payment_account(guild_id, payment_method, account_info):
    conn = get_db_connection()
    conn.execute('''
        INSERT INTO guild_payment_accounts (guild_id, payment_method, account_info)
        VALUES (?, ?, ?)
        ON CONFLICT(guild_id, payment_method) DO UPDATE SET
            account_info = excluded.account_info
    ''', (str(guild_id), payment_method, account_info))
    conn.commit()
    conn.close()

def delete_guild_payment_account(guild_id, payment_method):
    conn = get_db_connection()
    conn.execute(
        'DELETE FROM guild_payment_accounts WHERE guild_id = ? AND payment_method = ?',
        (str(guild_id), payment_method),
    )
    deleted = conn.total_changes
    conn.commit()
    conn.close()
    return deleted > 0

def set_order_guild_id(order_num, guild_id):
    conn = get_db_connection()
    conn.execute(
        'UPDATE orders SET guild_id = ? WHERE order_num = ? AND guild_id IS NULL',
        (str(guild_id), order_num),
    )
    conn.commit()
    conn.close()

def try_start_order_review(order_num, guild_id):
    conn = get_db_connection()
    cursor = conn.execute(
        '''UPDATE orders SET review_submitted = 1
           WHERE order_num = ? AND guild_id = ?
             AND status = 'COMPLETED'
             AND COALESCE(review_submitted, 0) = 0''',
        (order_num, str(guild_id)),
    )
    conn.commit()
    updated = cursor.rowcount > 0
    conn.close()
    return updated

def cancel_order_review(order_num, guild_id):
    conn = get_db_connection()
    conn.execute(
        '''UPDATE orders SET review_submitted = 0
           WHERE order_num = ? AND guild_id = ?''',
        (order_num, str(guild_id)),
    )
    conn.commit()
    conn.close()

def get_pending_auto_reviews(cutoff_time):
    conn = get_db_connection()
    rows = conn.execute(
        '''SELECT * FROM orders
           WHERE status = 'COMPLETED'
             AND COALESCE(review_submitted, 0) = 0
             AND completed_time IS NOT NULL
             AND completed_time <= ?
           ORDER BY completed_time ASC''',
        (cutoff_time,),
    ).fetchall()
    conn.close()
    return rows

def get_pending_reviews():
    conn = get_db_connection()
    rows = conn.execute(
        '''SELECT * FROM orders
           WHERE status = 'COMPLETED'
             AND COALESCE(review_submitted, 0) = 0
             AND guild_id IS NOT NULL
           ORDER BY completed_time ASC'''
    ).fetchall()
    conn.close()
    return rows

def create_order(
    guild_id, customer_id, order_type, price, players_needed, remarks, admin_id,
    ticket_msg_id=None, ticket_channel_id=None,
):
    import random
    import string
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        conn.execute('BEGIN IMMEDIATE')

        # Reserve the write transaction while checking and creating the unique 8-character ID.
        while True:
            letters = ''.join(random.choices(string.ascii_uppercase, k=5))
            digits = ''.join(random.choices(string.digits, k=3))
            order_num = letters + digits
            existing = cursor.execute(
                'SELECT 1 FROM orders WHERE order_num = ?',
                (order_num,),
            ).fetchone()
            if not existing:
                break

        created_at = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cursor.execute('''
        INSERT INTO orders (
            guild_id, order_num, customer_id, type, price, players_needed,
            remarks, status, admin_id, created_at, player_ids,
            ticket_msg_id, ticket_channel_id
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, 'PENDING', ?, ?, '[]', ?, ?)
        ''', (
            str(guild_id), order_num, customer_id, order_type, price,
            players_needed, remarks, admin_id, created_at,
            str(ticket_msg_id) if ticket_msg_id else None,
            str(ticket_channel_id) if ticket_channel_id else None,
        ))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return order_num

def get_order(order_num, guild_id=None):
    conn = get_db_connection()
    if guild_id is None:
        order = conn.execute('SELECT * FROM orders WHERE order_num = ?', (order_num,)).fetchone()
    else:
        order = conn.execute(
            'SELECT * FROM orders WHERE order_num = ? AND guild_id = ?',
            (order_num, str(guild_id)),
        ).fetchone()
    conn.close()
    return order

def update_order_payment(order_num, payment_method, guild_id=None, expected_status=None):
    conn = get_db_connection()
    payment_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    query = (
        'UPDATE orders SET payment_method = ?, payment_time = ? '
        'WHERE order_num = ?'
    )
    params = (payment_method, payment_time, order_num)
    if guild_id is not None:
        query += ' AND guild_id = ?'
        params += (str(guild_id),)
    if expected_status is not None:
        query += ' AND status = ?'
        params += (expected_status,)
    cursor = conn.execute(query, params)
    conn.commit()
    conn.close()
    return cursor.rowcount > 0

def update_order_status(order_num, status, guild_id=None, expected_status=None):
    conn = get_db_connection()
    if status == 'COMPLETED':
        completed_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        query = 'UPDATE orders SET status = ?, completed_time = ? WHERE order_num = ?'
        params = (status, completed_time, order_num)
    else:
        query = 'UPDATE orders SET status = ? WHERE order_num = ?'
        params = (status, order_num)
    if guild_id is not None:
        query += ' AND guild_id = ?'
        params += (str(guild_id),)
    if expected_status is not None:
        query += ' AND status = ?'
        params += (expected_status,)
    cursor = conn.execute(query, params)
    conn.commit()
    conn.close()
    return cursor.rowcount > 0

def add_player_to_order(order_num, player_id, guild_id=None):
    conn = get_db_connection()
    try:
        conn.execute('BEGIN IMMEDIATE')
        query = (
            "SELECT player_ids, players_needed FROM orders "
            "WHERE order_num = ? AND status = 'WAITING'"
        )
        params = [order_num]
        if guild_id is not None:
            query += ' AND guild_id = ?'
            params.append(str(guild_id))
        order = conn.execute(query, params).fetchone()
        if not order:
            return False, []

        raw_players = order['player_ids']
        try:
            player_ids = json.loads(raw_players) if raw_players else []
        except (json.JSONDecodeError, TypeError):
            player_ids = []

        if player_id in player_ids or len(player_ids) >= order['players_needed']:
            return False, player_ids

        player_ids.append(player_id)
        new_status = (
            'IN_PROGRESS'
            if len(player_ids) >= order['players_needed']
            else 'WAITING'
        )
        query = 'UPDATE orders SET player_ids = ?, status = ? WHERE order_num = ?'
        params = [json.dumps(player_ids), new_status, order_num]
        if guild_id is not None:
            query += ' AND guild_id = ?'
            params.append(str(guild_id))
        conn.execute(query, params)
        conn.commit()
        return True, player_ids
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

def get_active_orders(guild_id=None):
    conn = get_db_connection()
    query = "SELECT * FROM orders WHERE status != 'COMPLETED' AND status != 'CANCELLED'"
    params = []
    if guild_id is not None:
        query += " AND guild_id = ?"
        params.append(str(guild_id))
    orders = conn.execute(query, params).fetchall()
    conn.close()
    return orders

def update_order_message(order_num, msg_id, channel_id, guild_id=None):
    conn = get_db_connection()
    query = 'UPDATE orders SET dispatch_msg_id = ?, dispatch_channel_id = ? WHERE order_num = ?'
    params = [str(msg_id), str(channel_id), order_num]
    if guild_id is not None:
        query += ' AND guild_id = ?'
        params.append(str(guild_id))
    conn.execute(query, params)
    conn.commit()
    conn.close()

def update_order_ticket_message(order_num, msg_id, channel_id, guild_id=None):
    conn = get_db_connection()
    query = (
        'UPDATE orders SET ticket_msg_id = ?, ticket_channel_id = ? '
        'WHERE order_num = ?'
    )
    params = [str(msg_id), str(channel_id), order_num]
    if guild_id is not None:
        query += ' AND guild_id = ?'
        params.append(str(guild_id))
    conn.execute(query, params)
    conn.commit()
    conn.close()

def get_shipped_stock_keys(channel_id, guild_id=None):
    """回傳指定庫存頻道中已出貨資料的訊息與行號。"""
    conn = get_db_connection()
    rows = conn.execute(
        'SELECT message_id, line_number FROM stock_shipments WHERE channel_id = ?'
        + (' AND guild_id = ?' if guild_id is not None else ''),
        (str(channel_id),) if guild_id is None else (str(channel_id), str(guild_id))
    ).fetchall()
    conn.close()
    return {(str(row['message_id']), row['line_number']) for row in rows}

def get_shipped_stock_accounts(product_name, guild_id=None):
    """回傳指定商品所有已出貨帳號，避免刪除後重貼時重複出貨。"""
    conn = get_db_connection()
    rows = conn.execute(
        'SELECT account FROM stock_shipments WHERE product_name = ?'
        + (' AND guild_id = ?' if guild_id is not None else ''),
        (product_name,) if guild_id is None else (product_name, str(guild_id))
    ).fetchall()
    conn.close()
    return {row['account'] for row in rows}

def record_stock_shipments(guild_id, product_name, channel_id, stock_items, shipped_by):
    """將出貨項目寫入資料庫，並回傳實際成功記錄的筆數。"""
    if not stock_items:
        return 0

    shipped_at = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_db_connection()
    cursor = conn.cursor()
    inserted = 0
    try:
        conn.execute('BEGIN IMMEDIATE')
        for item in stock_items:
            cursor.execute(
                '''INSERT OR IGNORE INTO stock_shipments
                   (guild_id, product_name, channel_id, message_id, line_number, account, password, shipped_by, shipped_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                (
                    str(guild_id), product_name, str(channel_id), str(item['message_id']), item['line_number'],
                    item['account'], item['password'], str(shipped_by), shipped_at
                )
            )
            inserted += cursor.rowcount
        if inserted != len(stock_items):
            conn.rollback()
            conn.close()
            return 0
        conn.commit()
    except sqlite3.Error:
        conn.rollback()
        conn.close()
        raise
    conn.close()
    return inserted

def export_orders_to_excel():
    conn = get_db_connection()
    orders = conn.execute('SELECT * FROM orders').fetchall()
    conn.close()

    if os.path.exists(EXCEL_NAME):
        wb = load_workbook(EXCEL_NAME)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["訂單編號", "老闆ID", "類型", "金額", "打手數量", "備註", "狀態", "付款方式", "建立時間", "付款時間", "完成時間"])

    # 清除舊數據（保留標題）
    ws.delete_rows(2, ws.max_row)

    for order in orders:
        ws.append([
            order['order_num'],
            order['customer_id'],
            order['type'],
            order['price'],
            order['players_needed'],
            order['remarks'],
            order['status'],
            order['payment_method'],
            order['created_at'],
            order['payment_time'],
            order['completed_time']
        ])

    wb.save(EXCEL_NAME)
    print(f"📊 已更新 Excel 紀錄: {EXCEL_NAME}")

def add_spending_adjustment(guild_id, customer_id, amount, reason, admin_id):
    conn = get_db_connection()
    cursor = conn.cursor()
    created_at = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cursor.execute('''
    INSERT INTO spending_adjustments (guild_id, customer_id, amount, reason, admin_id, created_at)
    VALUES (?, ?, ?, ?, ?, ?)
    ''', (str(guild_id), str(customer_id), float(amount), reason or "管理員手動調整", str(admin_id), created_at))
    conn.commit()
    conn.close()

def get_spending_adjustments(guild_id, customer_id, limit=5):
    conn = get_db_connection()
    adjustments = conn.execute(
        'SELECT * FROM spending_adjustments WHERE guild_id = ? AND customer_id = ? ORDER BY id DESC LIMIT ?',
        (str(guild_id), str(customer_id), limit)
    ).fetchall()
    conn.close()
    return adjustments

def get_user_spending(guild_id, customer_id):
    conn = get_db_connection()
    cid = str(customer_id)
    current_month_prefix = datetime.datetime.now().strftime("%Y-%m")
    
    # 1. 已結單訂單總金額與數量
    row_total_orders = conn.execute(
        "SELECT SUM(price) as total_price, COUNT(*) as count FROM orders WHERE guild_id = ? AND customer_id = ? AND status = 'COMPLETED'",
        (str(guild_id), cid)
    ).fetchone()
    total_orders_price = row_total_orders['total_price'] if row_total_orders and row_total_orders['total_price'] is not None else 0.0
    completed_count = row_total_orders['count'] if row_total_orders else 0

    # 2. 當月已結單訂單金額 (優先判定 completed_time，無則用 created_at)
    row_monthly_orders = conn.execute(
        """SELECT SUM(price) as monthly_price FROM orders 
           WHERE guild_id = ? AND customer_id = ? AND status = 'COMPLETED' 
           AND (completed_time LIKE ? OR (completed_time IS NULL AND created_at LIKE ?))""",
        (str(guild_id), cid, f"{current_month_prefix}%", f"{current_month_prefix}%")
    ).fetchone()
    monthly_orders_price = row_monthly_orders['monthly_price'] if row_monthly_orders and row_monthly_orders['monthly_price'] is not None else 0.0

    # 3. 手動調整金額 (總計與當月)
    row_total_adj = conn.execute(
        "SELECT SUM(amount) as total_adj FROM spending_adjustments WHERE guild_id = ? AND customer_id = ?",
        (str(guild_id), cid)
    ).fetchone()
    total_adj = row_total_adj['total_adj'] if row_total_adj and row_total_adj['total_adj'] is not None else 0.0

    row_monthly_adj = conn.execute(
        "SELECT SUM(amount) as monthly_adj FROM spending_adjustments WHERE guild_id = ? AND customer_id = ? AND created_at LIKE ?",
        (str(guild_id), cid, f"{current_month_prefix}%")
    ).fetchone()
    monthly_adj = row_monthly_adj['monthly_adj'] if row_monthly_adj and row_monthly_adj['monthly_adj'] is not None else 0.0

    conn.close()

    total_spending = total_orders_price + total_adj
    monthly_spending = monthly_orders_price + monthly_adj

    return {
        "customer_id": cid,
        "total_spending": total_spending,
        "monthly_spending": monthly_spending,
        "total_orders_price": total_orders_price,
        "monthly_orders_price": monthly_orders_price,
        "total_adj": total_adj,
        "monthly_adj": monthly_adj,
        "completed_count": completed_count
    }
